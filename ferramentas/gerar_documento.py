"""Gera o documento ABNT (DOCX e PDF) a partir das specs e do modelo da disciplina.

Fontes lidas
  specs/*/spec.md                       requisitos RF, RNF e RN
  docs/casos_de_uso.md                  descricoes dos casos de uso
  docs/entrega/conteudo_documento.md    texto das secoes
  docs/entrega/capa.md                  dados da capa
  docs/uml/*.png                        diagramas
  docs/entrega/modelo/*.docx            modelo ABNT da disciplina

Saidas
  docs/entrega/Collection_Social_Documentacao.docx
  docs/entrega/Collection_Social_Documentacao.pdf
  docs/rastreabilidade.md
"""

import copy
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

RAIZ = Path(__file__).resolve().parent.parent
ENTREGA = RAIZ / "docs" / "entrega"
MODELO = ENTREGA / "modelo" / "modelo_projeto_abnt_comentado.docx"
CONTEUDO = ENTREGA / "conteudo_documento.md"
CAPA = ENTREGA / "capa.md"
CASOS_DE_USO = RAIZ / "docs" / "casos_de_uso.md"
RASTREABILIDADE = RAIZ / "docs" / "rastreabilidade.md"
SAIDA_DOCX = ENTREGA / "Collection_Social_Documentacao.docx"
SAIDA_PDF = ENTREGA / "Collection_Social_Documentacao.pdf"

FONTE = "Arial"
PRETO = "000000"
VERMELHO = "EE0000"
NUM_TITULOS = "17"
NUM_MARCADORES = "10"
FONTE_LEGENDA = "Fonte. Elaborado pelos autores (2026)."

# A4 com margens ABNT (NBR 14724), em twips
A4_LARGURA, A4_ALTURA = 11906, 16838
MARGEM_SUP_ESQ, MARGEM_INF_DIR = 1701, 1134
LARGURA_UTIL_CM = 16.0
ALTURA_UTIL_CM = 21.5
LARGURA_UTIL_PAISAGEM_CM = 24.7
ALTURA_UTIL_PAISAGEM_CM = 13.5

PADROES_CAPA = {
    "Curso": "NOME DO CURSO",
    "Apresentação": "Link da apresentação",
    "Integrantes": "Nome Completo RGM",
}


# ---------------------------------------------------------------- leitura


def ler_capa():
    dados, integrantes, em_integrantes = {}, [], False
    for linha in CAPA.read_text(encoding="utf8").splitlines():
        sub = re.match(r"^\s{2,}\* (.+)$", linha)
        item = re.match(r"^\* ([^.]+)\.\s*(.*)$", linha)
        if sub and em_integrantes:
            integrantes.append(sub.group(1).strip())
        elif item:
            chave, valor = item.group(1).strip(), item.group(2).strip()
            em_integrantes = chave == "Integrantes"
            if not em_integrantes:
                dados[chave] = valor
    dados["Integrantes"] = integrantes
    return dados


def ler_requisitos():
    requisitos = {"RF": [], "RNF": [], "RN": []}
    padrao = re.compile(r"^\* ((RNF|RF|RN)(\d+)) (.+)$")
    for spec in sorted(RAIZ.glob("specs/*/spec.md")):
        for linha in spec.read_text(encoding="utf8").splitlines():
            m = padrao.match(linha)
            if m:
                requisitos[m.group(2)].append((int(m.group(3)), m.group(1), m.group(4)))
    for tipo in requisitos:
        requisitos[tipo].sort()
        numeros = [n for n, _, _ in requisitos[tipo]]
        esperado = list(range(1, len(numeros) + 1))
        if numeros != esperado:
            sys.exit(f"Numeração de {tipo} com falhas ou repetições {numeros}")
    return requisitos


def ler_casos_de_uso():
    casos, atual, campo = [], None, None
    for linha in CASOS_DE_USO.read_text(encoding="utf8").splitlines():
        cabecalho = re.match(r"^## (UC\d+) (.+)$", linha)
        item = re.match(r"^\* ([^.]+)\.\s*(.*)$", linha)
        passo = re.match(r"^\s{2,}(\d+)\. (.+)$", linha)
        alternativo = re.match(r"^\s{2,}\* (.+)$", linha)
        if cabecalho:
            atual = {"id": cabecalho.group(1), "nome": cabecalho.group(2), "principal": [], "alternativo": []}
            casos.append(atual)
        elif atual is None:
            continue
        elif passo:
            atual["principal"].append(f"{passo.group(1)}. {passo.group(2)}")
        elif alternativo:
            atual["alternativo"].append(alternativo.group(1))
        elif item:
            campo = item.group(1).strip()
            atual[campo] = item.group(2).strip()
    return casos


def ids(texto):
    return re.findall(r"\b(?:RNF|RF|RN)\d+\b", texto or "")


def validar_rastreabilidade(requisitos, casos):
    rf_existentes = {i for _, i, _ in requisitos["RF"]}
    rn_existentes = {i for _, i, _ in requisitos["RN"]}
    cobertos = set()
    for caso in casos:
        for ref in ids(caso.get("Requisitos")):
            if ref not in rf_existentes:
                sys.exit(f"{caso['id']} cita {ref}, que não existe nas specs")
            cobertos.add(ref)
        for ref in ids(caso.get("Regras")):
            if ref not in rn_existentes:
                sys.exit(f"{caso['id']} cita {ref}, que não existe nas specs")
    faltando = sorted(rf_existentes - cobertos, key=lambda r: int(r[2:]))
    if faltando:
        sys.exit(f"Requisitos sem caso de uso {faltando}")


def ler_conteudo():
    blocos, paragrafo = [], []

    def fechar():
        if paragrafo:
            blocos.append(("paragrafo", " ".join(paragrafo)))
            paragrafo.clear()

    for linha in CONTEUDO.read_text(encoding="utf8").splitlines():
        titulo = re.match(r"^(#{1,3}) (.+)$", linha)
        figura = re.match(r"^!\[(.+)\]\((.+)\)(\{paisagem\})?$", linha)
        marcador = re.match(r"^\{\{(\w+)\}\}$", linha)
        lista = re.match(r"^\* (.+)$", linha)
        if titulo:
            fechar()
            blocos.append(("titulo", len(titulo.group(1)), titulo.group(2)))
        elif figura:
            fechar()
            caminho = (CONTEUDO.parent / figura.group(2)).resolve()
            blocos.append(("figura", figura.group(1), caminho, bool(figura.group(3))))
        elif marcador:
            fechar()
            blocos.append(("marcador", marcador.group(1)))
        elif lista:
            fechar()
            blocos.append(("lista", lista.group(1)))
        elif linha.strip():
            paragrafo.append(linha.strip())
        else:
            fechar()
    fechar()
    return blocos


# ---------------------------------------------------------------- XML


def elemento(tag, **atributos):
    el = OxmlElement(tag)
    for nome, valor in atributos.items():
        el.set(qn(nome.replace("_", ":", 1)), str(valor))
    return el


def propriedades_run(negrito=False, tamanho=24, cor=PRETO, italico=False):
    rpr = elemento("w:rPr")
    rpr.append(elemento("w:rFonts", w_ascii=FONTE, w_hAnsi=FONTE, w_cs=FONTE))
    if negrito:
        rpr.append(elemento("w:b"))
        rpr.append(elemento("w:bCs"))
    if italico:
        rpr.append(elemento("w:i"))
    rpr.append(elemento("w:color", w_val=cor))
    rpr.append(elemento("w:sz", w_val=tamanho))
    rpr.append(elemento("w:szCs", w_val=tamanho))
    return rpr


def run(texto, **formato):
    r = elemento("w:r")
    r.append(propriedades_run(**formato))
    t = elemento("w:t")
    t.text = texto
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    return r


def runs_com_negrito(texto, tamanho=24, cor=PRETO):
    partes = re.split(r"(\*\*[^*]+\*\*)", texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("**"):
            yield run(parte[2:-2], negrito=True, tamanho=tamanho, cor=cor)
        else:
            yield run(parte, tamanho=tamanho, cor=cor)


def paragrafo(texto="", alinhamento="both", recuo=True, entrelinha=360, depois=120, antes=0, tamanho=24, estilo=None):
    p = elemento("w:p")
    ppr = elemento("w:pPr")
    if estilo:
        ppr.append(elemento("w:pStyle", w_val=estilo))
    ppr.append(elemento("w:spacing", w_before=antes, w_after=depois, w_line=entrelinha, w_lineRule="auto"))
    if recuo:
        ppr.append(elemento("w:ind", w_firstLine=709))
    ppr.append(elemento("w:jc", w_val=alinhamento))
    p.append(ppr)
    for r in runs_com_negrito(texto, tamanho=tamanho):
        p.append(r)
    return p


def titulo(nivel, texto):
    p = elemento("w:p")
    ppr = elemento("w:pPr")
    ppr.append(elemento("w:pStyle", w_val="Ttulo1"))
    numpr = elemento("w:numPr")
    numpr.append(elemento("w:ilvl", w_val=nivel - 1))
    numpr.append(elemento("w:numId", w_val=NUM_TITULOS))
    ppr.append(numpr)
    ppr.append(elemento("w:spacing", w_before=360, w_after=240))
    ppr.append(elemento("w:outlineLvl", w_val=nivel - 1))
    ppr.append(propriedades_run(negrito=True))
    p.append(ppr)
    p.append(run(texto.upper() if nivel == 1 else texto, negrito=True))
    return p


def item_lista(texto):
    p = elemento("w:p")
    ppr = elemento("w:pPr")
    ppr.append(elemento("w:pStyle", w_val="PargrafodaLista"))
    numpr = elemento("w:numPr")
    numpr.append(elemento("w:ilvl", w_val=0))
    numpr.append(elemento("w:numId", w_val=NUM_MARCADORES))
    ppr.append(numpr)
    ppr.append(elemento("w:spacing", w_after=60, w_line=360, w_lineRule="auto"))
    ppr.append(elemento("w:jc", w_val="both"))
    p.append(ppr)
    for r in runs_com_negrito(texto):
        p.append(r)
    return p


def legenda(texto):
    p = paragrafo(texto, alinhamento="center", recuo=False, entrelinha=240, antes=240, depois=120, tamanho=20)
    p.find(qn("w:pPr")).insert(0, elemento("w:keepNext"))
    return p


def fonte_legenda():
    return paragrafo(FONTE_LEGENDA, alinhamento="center", recuo=False, entrelinha=240, depois=240, tamanho=20)


def definir_texto(p, texto, cor=PRETO, negrito=None):
    """Troca o texto de um paragrafo preservando a formatacao do primeiro run."""
    runs = p.findall(qn("w:r"))
    base = copy.deepcopy(runs[0].find(qn("w:rPr"))) if runs else propriedades_run()
    for r in runs:
        p.remove(r)
    antigo = base.find(qn("w:color"))
    if antigo is not None:
        base.remove(antigo)
    tamanho = base.find(qn("w:sz"))
    if tamanho is None:
        tamanho = base.find(qn("w:szCs"))
    if tamanho is not None:
        tamanho.addprevious(elemento("w:color", w_val=cor))
    else:
        base.append(elemento("w:color", w_val=cor))
    if negrito is not None:
        for tag in ("w:b", "w:bCs"):
            antigo = base.find(qn(tag))
            if antigo is not None:
                base.remove(antigo)
        if negrito:
            base.insert(1, elemento("w:b"))
    for i, linha in enumerate(texto.split("\n")):
        r = elemento("w:r")
        r.append(copy.deepcopy(base))
        if i:
            r.append(elemento("w:br"))
        t = elemento("w:t")
        t.text = linha
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        p.append(r)
    ppr_rpr = p.find(qn("w:pPr") + "/" + qn("w:rPr"))
    if ppr_rpr is not None:
        cor_ppr = ppr_rpr.find(qn("w:color"))
        if cor_ppr is not None:
            ppr_rpr.remove(cor_ppr)


def texto_de(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


# ---------------------------------------------------------------- secoes


def ajustar_pagina(sectpr, paisagem=False):
    pgsz = sectpr.find(qn("w:pgSz"))
    pgmar = sectpr.find(qn("w:pgMar"))
    if paisagem:
        pgsz.set(qn("w:w"), str(A4_ALTURA))
        pgsz.set(qn("w:h"), str(A4_LARGURA))
        pgsz.set(qn("w:orient"), "landscape")
    else:
        pgsz.set(qn("w:w"), str(A4_LARGURA))
        pgsz.set(qn("w:h"), str(A4_ALTURA))
        if pgsz.get(qn("w:orient")) is not None:
            del pgsz.attrib[qn("w:orient")]
    pgmar.set(qn("w:top"), str(MARGEM_SUP_ESQ))
    pgmar.set(qn("w:left"), str(MARGEM_SUP_ESQ))
    pgmar.set(qn("w:bottom"), str(MARGEM_INF_DIR))
    pgmar.set(qn("w:right"), str(MARGEM_INF_DIR))


def quebra_de_secao(base, paisagem):
    sectpr = copy.deepcopy(base)
    numeracao = sectpr.find(qn("w:pgNumType"))
    if numeracao is not None:
        sectpr.remove(numeracao)
    ajustar_pagina(sectpr, paisagem)
    p = elemento("w:p")
    ppr = elemento("w:pPr")
    ppr.append(elemento("w:spacing", w_after=0, w_line=240, w_lineRule="auto"))
    ppr.append(sectpr)
    p.append(ppr)
    return p


# ---------------------------------------------------------------- montagem


class Montador:
    def __init__(self, documento, tabela_modelo, sectpr_corpo):
        self.doc = documento
        self.tabela_modelo = tabela_modelo
        self.sectpr_corpo = sectpr_corpo
        self.elementos = []
        self.figuras = 0
        self.quadros = 0
        self.ultima_quebra_paisagem = None

    def figura(self, titulo_figura, caminho, paisagem):
        self.figuras += 1
        if paisagem and self.elementos and self.elementos[-1] is self.ultima_quebra_paisagem:
            # figuras em paisagem seguidas ficam na mesma secao, sem pagina em branco entre elas
            self.elementos.pop()
        elif paisagem:
            self.elementos.append(quebra_de_secao(self.sectpr_corpo, paisagem=False))
        largura_max = LARGURA_UTIL_PAISAGEM_CM if paisagem else LARGURA_UTIL_CM
        altura_max = ALTURA_UTIL_PAISAGEM_CM if paisagem else ALTURA_UTIL_CM
        with Image.open(caminho) as imagem:
            proporcao = imagem.height / imagem.width
        largura = min(largura_max, altura_max / proporcao)
        self.elementos.append(legenda(f"**Figura {self.figuras}.** {titulo_figura}"))
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(caminho), width=Cm(largura))
        self.elementos.append(p._p)
        self.elementos.append(fonte_legenda())
        if paisagem:
            self.ultima_quebra_paisagem = quebra_de_secao(self.sectpr_corpo, paisagem=True)
            self.elementos.append(self.ultima_quebra_paisagem)

    def requisitos(self, lista, tipo):
        for _, identificador, texto in lista:
            if tipo == "RN":
                m = re.match(r"^(\([^)]*\))\s*(.*)$", texto)
                destaque, resto = f"{identificador} {m.group(1)}", m.group(2)
            else:
                nome, _, resto = texto.partition(". ")
                destaque = f"{identificador} {nome}."
            self.elementos.append(paragrafo(f"**{destaque}** {resto}"))

    def tabela_quadro(self, titulo_quadro, linhas, larguras):
        self.quadros += 1
        self.elementos.append(legenda(f"**Quadro {self.quadros}.** {titulo_quadro}"))
        tabela = self.doc.add_table(rows=len(linhas), cols=len(larguras))
        tabela.style = self.doc.styles["Table Grid"]
        tabela.autofit = False
        for j, largura in enumerate(larguras):
            tabela.columns[j].width = Cm(largura)
        for i, valores in enumerate(linhas):
            for j, valor in enumerate(valores):
                celula = tabela.cell(i, j)
                celula.width = Cm(larguras[j])
                p = celula.paragraphs[0]._p
                p.append(elemento("w:pPr"))
                p.find(qn("w:pPr")).append(elemento("w:spacing", w_after=0, w_line=240, w_lineRule="auto"))
                p.append(run(valor, negrito=(i == 0), tamanho=20))
        self.elementos.append(tabela._tbl)
        self.elementos.append(fonte_legenda())

    def caso_de_uso(self, caso):
        self.quadros += 1
        self.elementos.append(legenda(f"**Quadro {self.quadros}.** {caso['id']} {caso['nome']}"))
        tabela = copy.deepcopy(self.tabela_modelo)
        linhas = tabela.findall(qn("w:tr"))
        extra = copy.deepcopy(linhas[-1])
        tabela.append(extra)
        linhas.append(extra)
        requisitos = " e ".join(", ".join(ids(caso["Requisitos"])).rsplit(", ", 1))
        regras = ids(caso.get("Regras"))
        rastreio = f"Requisitos {requisitos}."
        if regras:
            rastreio += " Regras " + " e ".join(", ".join(regras).rsplit(", ", 1)) + "."
        conteudo = [
            ("Caso de Uso", [f"{caso['id']} {caso['nome']}"]),
            ("Ator(es)", [caso["Atores"]]),
            ("Descrição", [caso["Descrição"]]),
            ("Precondição", [caso["Precondição"]]),
            ("Fluxo Principal", caso["principal"]),
            ("Fluxo Alternativo", caso["alternativo"] or ["Não há."]),
            ("Condições Finais", [caso["Condições finais"]]),
            ("Rastreabilidade", [rastreio]),
        ]
        for linha, (rotulo, valores) in zip(linhas, conteudo):
            celula_rotulo, celula_valor = linha.findall(qn("w:tc"))
            self._preencher_celula(celula_rotulo, [rotulo], negrito=True)
            self._preencher_celula(celula_valor, valores)
        self.elementos.append(tabela)
        self.elementos.append(fonte_legenda())

    @staticmethod
    def _preencher_celula(celula, linhas, negrito=False):
        for p in celula.findall(qn("w:p")):
            celula.remove(p)
        for texto in linhas:
            p = elemento("w:p")
            ppr = elemento("w:pPr")
            ppr.append(elemento("w:spacing", w_after=40, w_line=240, w_lineRule="auto"))
            p.append(ppr)
            p.append(run(texto, negrito=negrito, tamanho=20))
            celula.append(p)


def preencher_capa(corpo, capa):
    paragrafos = [el for el in corpo if el.tag == qn("w:p")]

    def achar(trecho):
        for p in paragrafos:
            if trecho in texto_de(p):
                return p
        sys.exit(f"Trecho da capa não encontrado {trecho}")

    def cor_de(chave, valor):
        return VERMELHO if PADROES_CAPA.get(chave) == valor else PRETO

    definir_texto(achar("CENTRO UNIVERSITÁRIO"), "UDF CENTRO UNIVERSITÁRIO DO DISTRITO FEDERAL")
    definir_texto(achar("NOME DO CURSO"), capa["Curso"], cor_de("Curso", capa["Curso"]))
    definir_texto(achar("DISCIPLINA"), f"DISCIPLINA {capa['Disciplina']}")
    definir_texto(achar("SISTEMA XYZ"), capa["Sistema"])
    definir_texto(achar("Alunos"), "Alunos")

    nomes = [p for p in paragrafos if "RGM" in texto_de(p)]
    integrantes = capa["Integrantes"] or [PADROES_CAPA["Integrantes"]]
    for p, nome in zip(nomes, integrantes):
        definir_texto(p, nome, cor_de("Integrantes", nome), negrito=False)
    for p in nomes[len(integrantes):]:
        corpo.remove(p)
    ultimo = nomes[min(len(nomes), len(integrantes)) - 1]
    for nome in integrantes[len(nomes):]:
        novo = copy.deepcopy(ultimo)
        definir_texto(novo, nome, cor_de("Integrantes", nome), negrito=False)
        ultimo.addnext(novo)
        ultimo = novo

    link = re.sub(r"^[a-z]+\W+", "", capa["Apresentação"]) if "//" in capa["Apresentação"] else capa["Apresentação"]
    definir_texto(achar("apresentação aqui"), link, cor_de("Apresentação", capa["Apresentação"]))
    definir_texto(achar("Professor"), "Professor", negrito=True)
    rodape = achar("Gabriel de Oliveira Alves")
    definir_texto(rodape, f"{capa['Professor']}\n\n\n{capa['Local']}\n{capa['Período']}")
    for r in rodape.findall(qn("w:r"))[3:]:
        r.find(qn("w:rPr")).insert(1, elemento("w:b"))


def montar():
    requisitos = ler_requisitos()
    casos = ler_casos_de_uso()
    validar_rastreabilidade(requisitos, casos)
    capa = ler_capa()
    blocos = ler_conteudo()

    documento = docx.Document(str(MODELO))
    corpo = documento.element.body
    elementos = list(corpo)

    fins_de_secao = [i for i, el in enumerate(elementos) if el.find(".//" + qn("w:sectPr")) is not None]
    inicio_corpo, fim_corpo = fins_de_secao[0], fins_de_secao[1]
    sectpr_corpo = elementos[fim_corpo].find(".//" + qn("w:sectPr"))
    tabela_modelo = copy.deepcopy(next(el for el in elementos[inicio_corpo:fim_corpo] if el.tag == qn("w:tbl")))

    preencher_capa(corpo, capa)

    for el in elementos[inicio_corpo + 1:fim_corpo]:
        corpo.remove(el)

    # ficha de autoavaliacao sem as orientacoes
    for el in elementos[fim_corpo + 1:]:
        if el.tag == qn("w:p") and el.find(".//" + qn("w:numPr")) is not None:
            corpo.remove(el)
    for tabela in documento.tables[-1:]:
        for linha in tabela.rows:
            for celula in linha.cells:
                for p in celula.paragraphs:
                    m = re.match(r"Nome completo 0(\d)", p.text)
                    if m and capa["Integrantes"] and int(m.group(1)) <= len(capa["Integrantes"]):
                        nome = capa["Integrantes"][int(m.group(1)) - 1]
                        if nome != PADROES_CAPA["Integrantes"]:
                            definir_texto(p._p, nome)

    montador = Montador(documento, tabela_modelo, sectpr_corpo)
    em_referencias = False
    for bloco in blocos:
        tipo = bloco[0]
        if tipo == "titulo":
            em_referencias = bloco[2] == "REFERÊNCIAS"
            montador.elementos.append(titulo(bloco[1], bloco[2]))
        elif tipo == "paragrafo" and em_referencias:
            montador.elementos.append(paragrafo(bloco[1], alinhamento="left", recuo=False, entrelinha=240, depois=240))
        elif tipo == "paragrafo":
            montador.elementos.append(paragrafo(bloco[1]))
        elif tipo == "lista":
            montador.elementos.append(item_lista(bloco[1]))
        elif tipo == "figura":
            montador.figura(bloco[1], bloco[2], bloco[3])
        elif bloco[1] in ("RF", "RNF", "RN"):
            montador.requisitos(requisitos[bloco[1]], bloco[1])
        elif bloco[1] == "CASOS_DE_USO":
            for caso in casos:
                montador.caso_de_uso(caso)
        elif bloco[1] == "RASTREABILIDADE":
            linhas = [("Caso de uso", "Requisitos funcionais", "Regras de negócio")]
            for caso in casos:
                regras = ", ".join(ids(caso.get("Regras"))) or "Nenhuma específica"
                linhas.append((f"{caso['id']} {caso['nome']}", ", ".join(ids(caso["Requisitos"])), regras))
            montador.tabela_quadro("Matriz de rastreabilidade entre casos de uso, requisitos e regras", linhas, [7.0, 4.5, 4.5])
        else:
            sys.exit(f"Marcador desconhecido {bloco[1]}")

    ancora = elementos[fim_corpo]
    for el in montador.elementos:
        ancora.addprevious(el)
    # o paragrafo de quebra de pagina do modelo ficava antes da secao, o fim da secao ja quebra
    for r in ancora.findall(qn("w:r")):
        ancora.remove(r)

    # paginas A4 com margens ABNT e numeracao iniciando no corpo
    secoes = [el for el in corpo.iter(qn("w:sectPr"))]
    for i, sectpr in enumerate(secoes):
        paisagem = sectpr.find(qn("w:pgSz")).get(qn("w:orient")) == "landscape"
        ajustar_pagina(sectpr, paisagem)
        numeracao = sectpr.find(qn("w:pgNumType"))
        if i == 1 and numeracao is None:
            antes = sectpr.find(qn("w:cols"))
            antes.addprevious(elemento("w:pgNumType", w_start=1))
        elif i > 1 and numeracao is not None:
            sectpr.remove(numeracao)

    # pede ao Word para atualizar o sumario ao abrir, na posicao exigida pelo esquema
    configuracoes = documento.settings.element
    posteriores = {"hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                   "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                   "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
                   "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"}
    atualizar = elemento("w:updateFields", w_val="true")
    seguinte = next((el for el in configuracoes if el.tag.split("}")[1] in posteriores), None)
    if seguinte is not None:
        seguinte.addprevious(atualizar)
    else:
        configuracoes.append(atualizar)

    documento.core_properties.title = "Collection Social"
    documento.core_properties.subject = "Documentação de negócios, requisitos e modelagem UML"
    documento.core_properties.author = "Grupo Collection Social"
    documento.save(str(SAIDA_DOCX))
    escrever_rastreabilidade(requisitos, casos)
    print(f"Gerado {SAIDA_DOCX.relative_to(RAIZ)} com {montador.figuras} figuras e {montador.quadros} quadros")


def escrever_rastreabilidade(requisitos, casos):
    linhas = [
        "# Matriz de Rastreabilidade",
        "",
        "Arquivo gerado por `ferramentas/gerar_documento.py` a partir das specs e de `docs/casos_de_uso.md`. Não edite à mão.",
        "",
        "## Casos de uso para requisitos e regras",
        "",
    ]
    for caso in casos:
        regras = ", ".join(ids(caso.get("Regras"))) or "nenhuma específica"
        linhas.append(f"* {caso['id']} {caso['nome']}. Requisitos {', '.join(ids(caso['Requisitos']))}. Regras {regras}.")
    linhas += ["", "## Requisitos funcionais para casos de uso", ""]
    for _, identificador, texto in requisitos["RF"]:
        usados = [c["id"] for c in casos if identificador in ids(c["Requisitos"])]
        linhas.append(f"* {identificador} {texto.partition('. ')[0]}. {', '.join(usados)}.")
    linhas += ["", "## Regras de negócio para casos de uso", ""]
    for _, identificador, texto in requisitos["RN"]:
        usados = [c["id"] for c in casos if identificador in ids(c.get("Regras"))]
        destino = ", ".join(usados) if usados else "validada nos testes das specs"
        linhas.append(f"* {identificador} {texto.split(')')[0]}). {destino}.")
    linhas += [
        "",
        "## Diagramas",
        "",
        "* Casos de uso. `docs/uml/01a_casos_de_uso_contas_acervo.puml` e `docs/uml/01b_casos_de_uso_social_trocas.puml`.",
        "* Classes. `docs/uml/02a_classes_contas_acervo.puml`, `docs/uml/02b_classes_social_moderacao.puml` e `docs/uml/02c_classes_trocas.puml`.",
        "* Atividades. `docs/uml/03_atividades_troca.puml`, processo de UC20 a UC23, regras RN25 a RN30.",
        "* Sequência. `docs/uml/04_sequencia_cadastrar_item.puml`, cenário do UC09, regras RN08 e RN09.",
        "* Estados. `docs/uml/05_estados_proposta.puml`, proposta de troca, regras RN26 a RN29.",
        "* Componentes. `docs/uml/06_componentes.puml`, arquitetura de `docs/arquitetura.md`.",
        "* Objetos. `docs/uml/07_objetos_troca.puml`, cenário do UC21, regras RN24 e RN25.",
        "",
    ]
    RASTREABILIDADE.write_text("\n".join(linhas), encoding="utf8")
    print(f"Gerado {RASTREABILIDADE.relative_to(RAIZ)}")


# ---------------------------------------------------------------- PDF

SCRIPT_UNO = r'''
import sys, time, uno
from com.sun.star.beans import PropertyValue

def prop(nome, valor):
    p = PropertyValue()
    p.Name = nome
    p.Value = valor
    return p

entrada, saida, porta = sys.argv[1], sys.argv[2], sys.argv[3]
local = uno.getComponentContext()
resolvedor = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
for tentativa in range(60):
    try:
        contexto = resolvedor.resolve("uno:socket,host=localhost,port=" + porta + ";urp;StarOffice.ComponentContext")
        break
    except Exception:
        time.sleep(0.5)
desktop = contexto.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", contexto)
doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(entrada), "_blank", 0, (prop("Hidden", True),))
for rodada in range(2):
    indices = doc.getDocumentIndexes()
    for i in range(indices.getCount()):
        indices.getByIndex(i).update()
    doc.refresh()
doc.storeToURL(uno.systemPathToFileUrl(saida), (prop("FilterName", "writer_pdf_Export"),))
doc.close(True)
try:
    desktop.terminate()
except Exception:
    pass
'''


def exportar_pdf():
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice não encontrado, PDF não gerado")
        return
    porta = "2083"
    with tempfile.TemporaryDirectory() as tmp:
        perfil = Path(tmp) / "perfil"
        script = Path(tmp) / "exportar.py"
        script.write_text(SCRIPT_UNO, encoding="utf8")
        processo = subprocess.Popen(
            [
                soffice,
                "--headless",
                "--invisible",
                "--norestore",
                f"-env:UserInstallation={perfil.as_uri()}",
                f"--accept=socket,host=localhost,port={porta};urp;",
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        try:
            subprocess.run([sys.executable, str(script), str(SAIDA_DOCX), str(SAIDA_PDF), porta], check=True, timeout=300)
        finally:
            time.sleep(1)
            processo.terminate()
            try:
                processo.wait(timeout=20)
            except subprocess.TimeoutExpired:
                processo.kill()
    print(f"Gerado {SAIDA_PDF.relative_to(RAIZ)}")


if __name__ == "__main__":
    montar()
    if "sem_pdf" not in sys.argv:
        exportar_pdf()
